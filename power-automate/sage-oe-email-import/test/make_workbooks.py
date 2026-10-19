"""Builds the two workbooks in this folder's parent:

  Sage_OE_Import_Template.xlsx - the empty template the flow copies for every email
  samples/OE_Import_SAMPLE.xlsx - what the script writes for samples/order-email.html

Run run-tests.sh first; it writes build/sample-output.json from the real script.
"""
import datetime
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

# Every column the script can write, in the order Sage should apply them.
TEMPLATE = {
    "Orders": ["ORDUNIQ", "ORDNUMBER", "CUSTOMER", "SHIPTO", "PONUMBER", "LOCATION",
               "ORDDATE", "EXPDATE", "REFERENCE", "DESC", "COMMENT"],
    "Order_Details": ["ORDUNIQ", "LINENUM", "LINETYPE", "ITEM", "QTYORDERED", "ORDUNIT", "UNITPRICE"],
}

HEADER_FONT = Font(name="Arial", size=10, bold=True)
BODY_FONT = Font(name="Arial", size=10)
EXCEL_EPOCH = datetime.date(1899, 12, 30)


def build(sheets, path):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows, formats in sheets:
        ws = wb.create_sheet(name)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                fmt = formats[r - 1][c - 1] if formats else ""
                if fmt and isinstance(value, (int, float)):
                    value = EXCEL_EPOCH + datetime.timedelta(days=value)
                cell = ws.cell(row=r, column=c, value=value if value != "" else None)
                cell.font = HEADER_FONT if r == 1 else BODY_FONT
                if fmt:
                    cell.number_format = fmt
        for c in range(1, len(rows[0]) + 1):
            width = max(len(str(row[c - 1])) for row in rows if c - 1 < len(row))
            ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 10), 60)
        ws.freeze_panes = "A2"
    wb.save(path)
    print("wrote", os.path.relpath(path, ROOT))


build([(name, [fields], None) for name, fields in TEMPLATE.items()],
      os.path.join(ROOT, "Sage_OE_Import_Template.xlsx"))

sample_json = os.path.join(HERE, "build", "sample-output.json")
if not os.path.exists(sample_json):
    sys.exit("build/sample-output.json missing - run ./run-tests.sh first")
with open(sample_json) as f:
    sample = json.load(f)
# Dates were written as Excel serials with a date format; formats mark which cells.
build([(s["name"], s["rows"], s["formats"]) for s in sample],
      os.path.join(ROOT, "samples", "OE_Import_SAMPLE.xlsx"))
